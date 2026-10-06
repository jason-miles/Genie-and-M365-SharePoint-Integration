#!/usr/bin/env python3
"""Generate a small synthetic document corpus (PDF + DOCX) in two "sites" for the downstream smoke test.

    pip install reportlab python-docx
    python3 tests/make_synthetic_corpus.py ./synthetic_corpus [--revised]

  SITE_POLICIES/  mri_preauthorisation_policy.pdf  dependants_cover_2025.pdf  provider_contracts_register.pdf
                  claims_handbook_long.pdf         hospital_admission_guide.docx
  SITE_HR/        salary_bands_2026.docx           disciplinary_procedure.pdf      (treated as RESTRICTED)

--revised produces the "change" scenario: the MRI policy changes (R6,500 -> R8,000), the hospital guide is omitted
(simulating a deletion) and an unparseable file corrupt_scan.pdf is added.
"""
import argparse
import pathlib

import docx
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def pdf(path, title, paras):
    c = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, title)
    y -= 36
    c.setFont("Helvetica", 11)
    for p in paras:
        for line in [p[i:i + 95] for i in range(0, len(p), 95)]:
            if y < 60:
                c.showPage()
                y = 800
                c.setFont("Helvetica", 11)
            c.drawString(50, y, line)
            y -= 16
        y -= 10
    c.save()


def word(path, heading, paras):
    d = docx.Document()
    d.add_heading(heading, 1)
    for p in paras:
        d.add_paragraph(p)
    d.save(str(path))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--revised", action="store_true")
    a = ap.parse_args()
    pol = pathlib.Path(a.out, "SITE_POLICIES")
    hr = pathlib.Path(a.out, "SITE_HR")
    pol.mkdir(parents=True, exist_ok=True)
    hr.mkdir(parents=True, exist_ok=True)

    if a.revised:
        pdf(pol / "mri_preauthorisation_policy.pdf", "MRI Pre-Authorisation Policy v4", [
            "Rule 1 (revised): An MRI scan requires pre-authorisation when the estimated cost exceeds R8,000.",
            "Rule 2: Requests must be submitted at least 72 hours before the scheduled scan."])
        (pol / "corrupt_scan.pdf").write_text("this is not a real pdf, just text pretending to be one, long enough to look like a file")
    else:
        pdf(pol / "mri_preauthorisation_policy.pdf", "MRI Pre-Authorisation Policy v3", [
            "Scope: This policy applies to all outpatient MRI scans requested under the Gold and Platinum plans.",
            "Rule 1: An MRI scan requires pre-authorisation when the estimated cost exceeds R6,500.",
            "Rule 2: Pre-authorisation requests must be submitted at least 48 hours before the scheduled scan, except in emergencies.",
            "Rule 3: Spinal MRI scans require a referral from a specialist and a documented failed six-week conservative treatment.",
            "Approval is valid for 30 days from the date of issue. Expired approvals must be re-submitted."])
        word(pol / "hospital_admission_guide.docx", "Hospital Admission Guide", [
            "Emergency admissions do not require pre-authorisation but must be notified within 24 hours of admission.",
            "Planned admissions require pre-authorisation 5 working days before the admission date."])

    pdf(pol / "dependants_cover_2025.pdf", "Dependants Cover Rules 2025", [
        "Child dependants are covered up to the age of 21, or 26 if a full-time student with proof of enrolment.",
        "Adult dependants must be declared at registration. Late additions carry a 3-month waiting period.",
        "The 2025 revision removed the previous 24-month waiting period for newborn dependants added within 30 days of birth."])
    pdf(pol / "provider_contracts_register.pdf", "Provider Contract Register", [
        "Contract 1: Alpha Radiology Group, expires 31 December 2026.",
        "Contract 2: Beta Oncology Partners, expires 30 June 2027.",
        "Contract 3: Gamma Pathology Labs, expires 31 December 2026.",
        "Contract 4: Delta Physiotherapy Network, expires 15 March 2028."])
    pdf(pol / "claims_handbook_long.pdf", "Claims Handbook (long)",
        ["Clause %d: Members must retain original invoices for %d months for audit purposes. " % (i, 12 + i)
         for i in range(1, 61)])
    word(hr / "salary_bands_2026.docx", "Salary Bands 2026 (RESTRICTED)", [
        "Band E salary range is R780,000 to R1,050,000 per year. This document is restricted to the HR team."])
    pdf(hr / "disciplinary_procedure.pdf", "Disciplinary Procedure (RESTRICTED)", [
        "First written warning remains valid for 6 months. Final written warning remains valid for 12 months.",
        "This document is restricted to HR and line managers."])
    print(f"wrote corpus to {a.out}{' (revised)' if a.revised else ''}")


if __name__ == "__main__":
    main()
